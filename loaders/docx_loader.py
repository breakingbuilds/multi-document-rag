"""
loaders/docx_loader.py
----------------------
Extracts text from Word (.docx) files, grouped by heading.

Why group by heading?
    A DOCX has no "pages" we can rely on (page breaks depend on the renderer),
    but it DOES have a heading hierarchy. Grouping paragraphs under the
    nearest heading gives each `LoadedDocument` a meaningful `section`
    locator such as "Leave at a Glance", which is far more useful in a
    citation than a byte offset.

Tables:
    python-docx exposes tables separately from paragraphs, but in the
    document body they are interleaved. We walk the underlying XML body in
    order so tables land in the correct section, and we flatten every table
    row into "Header: value | Header: value" lines. That keeps tabular facts
    (e.g. "Annual leave | 20 paid days") searchable by the embedding model.
"""

from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from .base import LoadedDocument, base_metadata, normalize_whitespace


def _iter_body_elements(doc):
    """Yield Paragraph and Table objects in true document order.

    python-docx keeps `doc.paragraphs` and `doc.tables` in separate lists,
    so we iterate the raw XML body and wrap each child in the matching
    python-docx class to preserve the original ordering.
    """
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]  # strip the XML namespace
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def _table_to_text(table: Table) -> str:
    """Flatten a table into one line per data row, prefixed by column names."""
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, *data = rows
    lines = []
    for row in data:
        pairs = [f"{h}: {v}" for h, v in zip(header, row) if v]
        lines.append(" | ".join(pairs))
    return "\n".join(lines)


def load_docx(path: Path) -> list[LoadedDocument]:
    """Return one LoadedDocument per heading-delimited section."""
    doc = Document(str(path))
    documents: list[LoadedDocument] = []

    current_heading = "Introduction"
    buffer: list[str] = []

    def flush():
        """Emit the accumulated section (if any) as a document."""
        if buffer and buffer[0].strip() == current_heading.strip() and not any(l.strip() for l in buffer[1:]):
            return  # heading with no body yet: merge it into the next section
        text = normalize_whitespace("\n".join(buffer))
        if text:
            metadata = base_metadata(path)
            metadata["section"] = current_heading
            documents.append(LoadedDocument(text=text, metadata=metadata))
        buffer.clear()

    for element in _iter_body_elements(doc):
        if isinstance(element, Paragraph):
            style = (element.style.name or "").lower()
            text = element.text.strip()
            if not text:
                continue
            if style.startswith("heading") or style == "title":
                # A new heading starts a new section: flush the previous one.
                flush()
                current_heading = text
                # Keep the heading inside the text too -- it is a strong
                # retrieval signal ("Resignation and Notice Period").
                buffer.append(text)
            elif "list" in style:
                buffer.append(f"- {text}")
            else:
                buffer.append(text)
        else:  # Table
            table_text = _table_to_text(element)
            if table_text:
                buffer.append(table_text)

    flush()  # don't forget the last section
    return documents
